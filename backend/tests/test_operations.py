"""Real processing, revision, sharing and API-key integration exercises."""

from datetime import UTC, datetime, timedelta
from uuid import UUID, uuid4

import pymupdf
import pytest
from conftest import ApiClient, make_pdf
from test_accounts import register

from transler.worker import perform_job

pytestmark = pytest.mark.integration


def upload(client: ApiClient, name: str, content: bytes) -> dict[str, str]:
    response = client.post(
        "/v1/documents", files={"file": (name, content, "application/octet-stream")}
    )
    assert response.status_code == 201, response.text
    return response.json()


def execute(
    client: ApiClient, document: dict[str, str], artifact: str, options: dict[str, object]
) -> dict[str, object]:
    response = client.post(
        "/v1/operations",
        headers={"Idempotency-Key": str(uuid4())},
        json={"document_id": document["id"], "input_artifact_id": artifact, "options": options},
    )
    assert response.status_code == 202, response.text
    perform_job(client.settings, UUID(response.json()["id"]))
    result = client.get("/v1/jobs/" + response.json()["id"]).json()
    assert result["status"] == "COMPLETED", result
    return result


def test_conversion_edit_and_revision_conflict(client: ApiClient) -> None:
    document = upload(client, "source.pdf", make_pdf(2))
    text = execute(
        client, document, document["artifact_id"], {"type": "conversion", "target_format": "txt"}
    )
    downloaded = client.get(f"/v1/artifacts/{text['artifact_id']}/download")
    assert downloaded.headers["content-type"].startswith("text/plain")
    assert "Welcome to Transler" in downloaded.text
    operation = {
        "type": "rotate_page",
        "page": 0,
        "rect": None,
        "destination": None,
        "text": None,
        "font_size": None,
        "angle": 90,
        "order": None,
        "source_language": None,
        "target_language": None,
    }
    options = {"type": "edit", "base_revision": 0, "operations": [operation]}
    edit = execute(client, document, document["artifact_id"], options)
    with pymupdf.open(
        stream=client.get(f"/v1/artifacts/{edit['artifact_id']}/download").content, filetype="pdf"
    ) as pdf:
        assert pdf[0].rotation == 90
        assert pdf.page_count == 2
    revisions = client.get(f"/v1/documents/{document['id']}/revisions").json()
    assert revisions[0]["revision"] == 1
    conflict = client.post(
        "/v1/operations",
        headers={"Idempotency-Key": str(uuid4())},
        json={
            "document_id": document["id"],
            "input_artifact_id": document["artifact_id"],
            "options": options,
        },
    )
    assert conflict.status_code == 409, conflict.text
    assert conflict.json()["code"] == "REVISION_CONFLICT"


def test_pdf_tools_encrypted_password_and_unlock(client: ApiClient) -> None:
    document = upload(client, "source.pdf", make_pdf(2))
    password = "test-document-password"
    protected = execute(
        client,
        document,
        document["artifact_id"],
        {
            "type": "pdf_tool",
            "tool": "protect",
            "artifact_ids": [],
            "pages": [],
            "password": password,
        },
    )
    with pymupdf.open(
        stream=client.get(f"/v1/artifacts/{protected['artifact_id']}/download").content,
        filetype="pdf",
    ) as pdf:
        assert pdf.is_encrypted
        assert pdf.authenticate(password)
    from sqlalchemy import text

    with client.database.connect() as connection:
        stored = connection.scalar(
            text("SELECT options->>'password' FROM jobs WHERE id=:id"), {"id": protected["id"]}
        )
        assert stored != password
    unlocked = execute(
        client,
        document,
        str(protected["artifact_id"]),
        {
            "type": "pdf_tool",
            "tool": "unlock",
            "artifact_ids": [],
            "pages": [],
            "password": password,
        },
    )
    with pymupdf.open(
        stream=client.get(f"/v1/artifacts/{unlocked['artifact_id']}/download").content,
        filetype="pdf",
    ) as pdf:
        assert not pdf.is_encrypted
        assert pdf.page_count == 2
    split = execute(
        client,
        document,
        str(unlocked["artifact_id"]),
        {"type": "pdf_tool", "tool": "split", "artifact_ids": [], "pages": [1], "password": None},
    )
    with pymupdf.open(
        stream=client.get(f"/v1/artifacts/{split['artifact_id']}/download").content, filetype="pdf"
    ) as pdf:
        assert pdf.page_count == 1
        assert "Page 2" in pdf[0].get_text()


def test_api_key_sharing_and_trash(client: ApiClient, pdf: bytes) -> None:
    document = upload(client, "source.pdf", pdf)
    register(client)
    expiry = (datetime.now(UTC) + timedelta(days=1)).isoformat()
    key = client.post("/v1/account/api-keys", json={"name": "Integration", "expires_at": expiry})
    assert key.status_code == 201, key.text
    import httpx

    with httpx.Client(base_url=str(client.base_url)) as external:
        response = external.get(
            "/v1/documents", headers={"Authorization": "Bearer " + key.json()["token"]}
        )
        assert response.status_code == 200, response.text
        assert response.json()[0]["id"] == document["id"]
        assert (
            external.get(
                "/v1/admin/overview", headers={"Authorization": "Bearer " + key.json()["token"]}
            ).status_code
            == 401
        )
        assert client.delete("/v1/account/api-keys/" + key.json()["key"]["id"]).status_code == 204
        assert (
            external.get(
                "/v1/documents", headers={"Authorization": "Bearer " + key.json()["token"]}
            ).status_code
            == 401
        )
        share = client.post(
            "/v1/shares", json={"artifact_id": document["artifact_id"], "expires_at": expiry}
        )
        assert share.status_code == 201, share.text
        path = "/v1/shared/" + share.json()["url"].rsplit("/", 1)[-1]
        assert external.get(path).content == pdf
        trashed = client.patch(
            f"/v1/documents/{document['id']}", json={"favorite": True, "trashed": True}
        )
        assert trashed.status_code == 200, trashed.text
        assert trashed.json()["favorite"]
        assert external.get(path).status_code == 404
        client.patch(
            f"/v1/documents/{document['id']}", json={"favorite": True, "trashed": False}
        ).raise_for_status()
        assert external.get(path).content == pdf
        assert client.delete("/v1/shares/" + share.json()["share"]["id"]).status_code == 204
        assert external.get(path).status_code == 404


@pytest.mark.translation
def test_real_advanced_pdf(client: ApiClient, pdf: bytes) -> None:
    document = upload(client, "source.pdf", pdf)
    result = execute(
        client,
        document,
        document["artifact_id"],
        {
            "type": "advanced",
            "source_language": "en",
            "target_language": "vi",
            "domain": "document translation",
            "custom_instructions": "",
            "glossary": [{"source": "scientific formulas", "target": "công thức khoa học"}],
        },
    )
    content = client.get(f"/v1/artifacts/{result['artifact_id']}/download").content
    with pymupdf.open(stream=content, filetype="pdf") as output:
        assert "dịch" in output[0].get_text().lower()
        assert "1" in output[0].get_text()
        output[0].get_pixmap(dpi=160).save("../.runtime/advanced-output.png")


@pytest.mark.translation
def test_real_markdown_translation_preserves_code(client: ApiClient) -> None:
    source = (
        b"# Research\n\nThe model uses 30 pages of training data.\n\n"
        b"```python\nresult = fit(X)\n```\n"
    )
    document = upload(client, "source.md", source)
    result = execute(
        client,
        document,
        document["artifact_id"],
        {
            "type": "format_translation",
            "source_language": "en",
            "target_language": "vi",
            "domain": "machine learning",
            "custom_instructions": "",
            "glossary": [{"source": "training data", "target": "dữ liệu huấn luyện"}],
        },
    )
    content = client.get(f"/v1/artifacts/{result['artifact_id']}/download").text
    assert "```python\nresult = fit(X)\n```" in content
    assert "30" in content
    assert "dữ liệu huấn luyện" in content


@pytest.mark.translation
def test_real_office_and_html_translation(client: ApiClient) -> None:
    from io import BytesIO

    from docx import Document
    from openpyxl import Workbook, load_workbook
    from pptx import Presentation
    from pptx.util import Inches

    word = Document()
    word.add_paragraph().add_run("Research result 30").bold = True
    word_bytes = BytesIO()
    word.save(word_bytes)
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "Research result 30"
    slide_bytes = BytesIO()
    slides.save(slide_bytes)
    workbook = Workbook()
    workbook.active["A1"] = "Research result 30"
    workbook.active["B1"] = "=1+2"
    book_bytes = BytesIO()
    workbook.save(book_bytes)
    inputs = (
        ("docx", word_bytes.getvalue()),
        ("pptx", slide_bytes.getvalue()),
        ("xlsx", book_bytes.getvalue()),
        (
            "html",
            b"<!doctype html><h1>Research result 30</h1><pre><b>fit(X)</b></pre>"
            b'<a href="https://example.com">Research result 30</a>',
        ),
        ("txt", b"Research result 30"),
    )
    for kind, source in inputs:
        document = upload(client, "research." + kind, source)
        result = execute(
            client,
            document,
            document["artifact_id"],
            {
                "type": "format_translation",
                "source_language": "en",
                "target_language": "vi",
                "domain": "research",
                "custom_instructions": "",
                "glossary": [{"source": "Research result", "target": "Kết quả nghiên cứu"}],
            },
        )
        content = client.get(f"/v1/artifacts/{result['artifact_id']}/download").content
        if kind == "docx":
            translated = Document(BytesIO(content))
            text = translated.paragraphs[0].text
            assert translated.paragraphs[0].runs[0].bold
        elif kind == "pptx":
            text = Presentation(BytesIO(content)).slides[0].shapes[0].text
        elif kind == "xlsx":
            translated_book = load_workbook(BytesIO(content))
            text = translated_book.active["A1"].value
            assert translated_book.active["B1"].value == "=1+2"
        else:
            text = content.decode()
            if kind == "html":
                assert (
                    "<pre><b>fit(X)</b></pre>" in text
                    and 'href="https://example.com"' in text
                    and "<!DOCTYPE html>" in text
                )
        assert "Kết quả nghiên cứu" in text and "30" in text, text
