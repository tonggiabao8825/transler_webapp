"""Exercise conversion and OCR against the real worker OS dependencies."""

import subprocess
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile


def verify_worker_adapters() -> None:
    import pymupdf
    from docx import Document
    from openpyxl import Workbook, load_workbook
    from pptx import Presentation
    from pptx.util import Inches

    from transler.config import LLMConfig
    from transler.conversion import convert
    from transler.editing import edit_pdf
    from transler.ocr import searchable_pdf
    from transler.processing_contracts import EditOptions

    with pymupdf.open() as pdf:
        page = pdf.new_page(width=500, height=500)
        page.insert_text((50, 50), "Research result 30", fontsize=18)
        for x in (50, 200, 350):
            page.draw_line((x, 100), (x, 180))
        for y in (100, 140, 180):
            page.draw_line((50, y), (350, y))
        for point, label in (
            ((60, 125), "Method"),
            ((210, 125), "Score"),
            ((60, 165), "Baseline"),
            ((210, 165), "98.5"),
        ):
            page.insert_text(point, label, fontsize=12)
        source = pdf.tobytes()
        image = page.get_pixmap(dpi=180).tobytes("png")
    for target in ("docx", "pptx", "xlsx", "png", "jpg", "html", "txt", "md"):
        result, suffix, _mime = convert(source, "pdf", target)
        assert result
        if target == "docx":
            assert Document(BytesIO(result)).paragraphs or Document(BytesIO(result)).tables
        elif target == "pptx":
            assert len(Presentation(BytesIO(result)).slides) == 1
        elif target == "xlsx":
            assert load_workbook(BytesIO(result)).active["A1"].value == "Method"
        elif target in {"png", "jpg"}:
            with ZipFile(BytesIO(result)) as archive:
                assert len(archive.namelist()) == 1
        print("PDF ->", target, "validated")
    doc = Document()
    doc.add_paragraph("Research result 30")
    word = BytesIO()
    doc.save(word)
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "Research result 30"
    presentation = BytesIO()
    slides.save(presentation)
    book = Workbook()
    book.active["A1"] = "Research result 30"
    workbook = BytesIO()
    book.save(workbook)
    for kind, data in (
        ("docx", word.getvalue()),
        ("pptx", presentation.getvalue()),
        ("xlsx", workbook.getvalue()),
        ("png", image),
        ("html", b'<h1>Research result 30</h1><a href="https://example.com">Link</a>'),
        ("txt", b"Research result 30"),
        ("md", b"# Research result 30"),
    ):
        result, _, _ = convert(data, kind, "pdf")
        with pymupdf.open(stream=result, filetype="pdf") as pdf:
            assert pdf.page_count >= 1
        print(kind, "-> PDF validated")
    scanned, _, _ = convert(image, "png", "pdf")
    recognized = searchable_pdf(scanned, "en")
    with pymupdf.open(stream=recognized, filetype="pdf") as pdf:
        assert "Research" in pdf[0].get_text()
        assert pdf[0].get_images()
        pdf[0].get_pixmap(dpi=160).save("/tmp/transler-ocr-verified.png")
    options = EditOptions.model_validate(
        {
            "type": "edit",
            "base_revision": 0,
            "operations": [
                {
                    "type": "replace_text",
                    "page": 0,
                    "rect": [45, 25, 400, 75],
                    "destination": None,
                    "text": "Updated result 30",
                    "font_size": 16,
                    "angle": None,
                    "order": None,
                    "source_language": None,
                    "target_language": None,
                }
            ],
        }
    )
    edited = edit_pdf(source, options, LLMConfig())
    with pymupdf.open(stream=edited, filetype="pdf") as pdf:
        assert "Updated result 30" in pdf[0].get_text()
        assert "Research result" not in pdf[0].get_text()
        pdf[0].get_pixmap(dpi=160).save("/tmp/transler-edit-verified.png")
    print("OCR and text replacement validated")


def test_worker_conversion_and_ocr_matrix() -> None:
    root = Path(__file__).resolve().parents[2]
    result = subprocess.run(
        ["docker", "compose", "exec", "-T", "worker", "python", "-"],
        input=Path(__file__).read_text() + "\nverify_worker_adapters()\n",
        capture_output=True,
        text=True,
        cwd=root,
        timeout=180,
        check=False,
    )
    assert result.returncode == 0, result.stdout + result.stderr
